# -*- coding: utf-8 -*-
"""
DAHITI + Hydroweb.next satellite altimetry - incremental update in four steps.

STEP 1  Check the merged station observations already saved
        (merged_altimetry_stations.csv): stations, observation count, last date.
STEP 2  Download new observations from the DAHITI and Hydroweb.next APIs
        - DAHITI: per station, only dates after the last saved one
          -> dahiti_water_levels_raw.xlsx
        - Hydroweb: rivers + lakes zips for the bbox (py_hydroweb)
          -> HYDROWEB_RIVERS_OPE.zip / HYDROWEB_LAKES_OPE.zip -> hydroweb_water_levels_raw.xlsx
STEP 3  Merge, keep only observations not merged yet, and for those:
          * running min / max / average / deviation, difference from previous
          * 15-day antecedent CHIRPS rainfall + ERA5-Land soil moisture,
            evapotranspiration, 2 m temperature, runoff (GEE, new rows only)
          * SRTM elevation (new stations only), up/downstream links,
            Lag Days over the last 12 observations
        -> appended to merged_altimetry_stations.csv / .xlsx
STEP 4  Upload the new station observations to GEE
        (projects/ee-penuelabi/assets/altimetry/merged_observations + stations)

Run in Google Colab. API keys: Colab secrets DAHITI_API_KEY / HYDROWEB_API_KEY
if set (key icon, left panel), otherwise the keys written in the CONFIGURATION section.
Station IDs are written as "source:id" in the link columns (e.g. "dahiti:213",
"hydroweb:12345") because DAHITI and Hydroweb IDs can collide.
"""

import os
import re
import sys
import math
import glob
import time
import shutil
import zipfile
import datetime
import subprocess
from io import StringIO

import numpy as np
import pandas as pd
import requests
import ee

# =========================================================================== #
# 1. CONFIGURATION                                                            #
# =========================================================================== #
PROJECT_ID = 'ee-penuelabi'

OUT_DIR = os.environ.get('ALTIMETRY_OUT_DIR', '/content/drive/MyDrive/colab_waterlevel/')
WORK_DIR = os.environ.get('ALTIMETRY_WORK_DIR', '/content/hydroweb_extract')
OUTPUT_BASENAME = 'merged_altimetry_stations'

# --- Download switches: set False to reuse what is already on Drive ---------
DOWNLOAD_DAHITI = True
DOWNLOAD_HYDROWEB = True

# --- Raw downloads (kept on Drive, updated each run) ------------------------
DAHITI_RAW_XLSX = os.path.join(OUT_DIR, 'dahiti_water_levels_raw.xlsx')
# Existing output of the old DAHITI script: read once to start from (never modified).
# Leave '' to pick the newest combined_water_levels_analyzed*.xlsx in OUT_DIR.
DAHITI_SEED_FILE = os.path.join(OUT_DIR, 'combined_water_levels_analyzed_20260922_093406.xlsx')
RIVERS_ZIP = os.path.join(OUT_DIR, 'HYDROWEB_RIVERS_OPE.zip')
LAKES_ZIP = os.path.join(OUT_DIR, 'HYDROWEB_LAKES_OPE.zip')
HYDROWEB_RAW_XLSX = os.path.join(OUT_DIR, 'hydroweb_water_levels_raw.xlsx')

# --- GEE caches so reruns only fetch what is new ----------------------------
CLIMATE_CACHE = os.path.join(OUT_DIR, 'cache_climate_15d.csv')
ELEV_CACHE = os.path.join(OUT_DIR, 'cache_elevation.csv')

START_YEAR = 2016            # observations before this are dropped at merge (both sources)
WINDOW_DAYS = 15             # antecedent window: 15 days ending ON the observation date
BUFFER_M = {'river': 5000, 'lake': 20000}   # climate sampling radius by type
ELEV_BUFFER_M = 1000
GEE_BATCH = 500              # features per getInfo call
RETRY_RECENT_DAYS = 60       # incomplete windows newer than this are re-requested next run

# Hydroweb download / filter box: [min_lon, min_lat, max_lon, max_lat]
HYDROWEB_BBOX = [23.2, -1.8, 36.1, 16.0]

# Onset-lag settings
MAX_LAG_DAYS = 60            # search lags 0..60 days (downstream responds later)
MAX_GAP_DAYS = 45            # don't interpolate across gaps longer than this
LAG_WINDOW_OBS = 12          # Lag Days uses the window covered by the station's last 12 observations
LAG_MIN_OVERLAP_DAYS = 60    # min overlapping daily values inside that window
MIN_R = 0.3                  # min correlation to accept a lag

# Extra links you want to add or override, incl. cross-source ones:
# [("hydroweb:12345", "dahiti:213"), ...]   (upstream, downstream)
MANUAL_LINKS = []

# --- STAGE 4: upload results to Earth Engine assets --------------------------
UPLOAD_TO_GEE = True
ASSET_FOLDER = f'projects/{PROJECT_ID}/assets/altimetry'
ROWS_PER_PART = 3000         # features sent per export request (keeps requests < 10 MB)

# --- Incremental outputs: fixed master files, new observations appended ------
MASTER_CSV = os.path.join(OUT_DIR, f'{OUTPUT_BASENAME}.csv')
MASTER_XLSX = os.path.join(OUT_DIR, f'{OUTPUT_BASENAME}.xlsx')
GEE_LEDGER = os.path.join(OUT_DIR, 'gee_uploaded_keys.csv')   # which rows are already in GEE
# Set True for ONE run to delete the master files + ledger and rebuild everything
# (e.g. after changing how a column is calculated). Set back to False afterwards.
REBUILD_OUTPUTS = False


# --- API keys: environment variable (GitHub secret) -> Colab secret -> key written here
try:
    from google.colab import userdata
    HYDROWEB_API_KEY = userdata.get("HYDROWEB_API_KEY")
except Exception:
    HYDROWEB_API_KEY = ""   # set as a GitHub secret

try:
    from google.colab import userdata
    DAHITI_API_KEY = userdata.get("DAHITI_API_KEY")
except Exception:
    DAHITI_API_KEY = ''   # set as a GitHub secret

HYDROWEB_API_KEY = (os.environ.get('HYDROWEB_API_KEY') or HYDROWEB_API_KEY or '').strip()
DAHITI_API_KEY = (os.environ.get('DAHITI_API_KEY') or DAHITI_API_KEY or '').strip()

URL_LIST = 'https://dahiti.dgfi.tum.de/api/v2/list-targets/'
URL_DOWNLOAD = 'https://dahiti.dgfi.tum.de/api/v2/download-water-level/'
DAHITI_LIST_ARGS = {'api_key': DAHITI_API_KEY, 'min_lon': 23.2, 'max_lon': 36.1,
                    'min_lat': -1.8, 'max_lat': 40.8}

DAHITI_IDS = [
    85, 213, 2686, 2687, 2688, 3246, 3247, 3764, 5959,
    8313, 11635, 11683, 11750, 11815, 11900, 15192, 15710, 16098,
    16099, 16985, 18043, 18741, 18743, 19747, 19749, 23333, 23334,
    23336, 23337, 23338, 23341, 23353, 23354, 23357, 23358, 23359, 34869, 38468,
    39376, 12258, 17741, 16686, 16323, 14196, 10179, 2, 2264, 11761, 12138, 12178,
    15556, 15557, 16846, 17743, 17749, 17750, 17752, 17753, 17754,
    38453, 38456, 967, 23344, 23343, 23342, 19752, 19751,
]

DAHITI_ORDER = """id\tafter\n2\t38471\n38471\t16846\n16846\t38456\n38456\t2264\n2264\t38454\n38454\t38455\n38455\t38453\n38453\t213\n213\t16099\n16099\t15192\n15192\t16098\n85\t16098\n16098\t967\n967\t15556\n15556\t11761\n11761\t17753\n17753\t17754\n17754\t3246\n3246\t17752\n3247\t17752\n17752\t17750\n17750\t17749\n17749\t15557\n2686\t39376\n39376\t2687\n23344\t23343\n23343\t19749\n19749\t23342\n23342\t23341\n23341\t23338\n18741\t23338\n2687\t2688\n2688\t23338\n23338\t23337\n23337\t23336\n23336\t15557\n15557\t17743\n19752\t17743\n12138\t19752\n17743\t19751\n19751\t12178\n12178\t10179\n10179\t14196\n14196\t16323\n16323\t16686\n16686\t17741\n17741\t12258\n23359\t23358\n23358\t23357\n23357\t18743\n18743\t18043\n18043\t11815\n11815\t10179\n23354\t23353\n23353\t3764\n3764\t19747\n19747\t23359\n11683\t19747\n23333\t11683\n23334\t23333\n11635\t23334\n11750\t11635"""

RAW_COLS = ['source', 'station_id', 'type', 'location', 'latitude', 'longitude',
            'river_key', 'date_dt', 'wse', 'uncertainty']
DAHITI_RAW_COLS = ['DAHITI-ID', 'Target Name', 'Target Type', 'Longitude', 'Latitude',
                   'date [yyyy-mm-dd]', 'water level', 'wse_u']

# =========================================================================== #
# 2. SETUP                                                                    #
# =========================================================================== #
try:
    from google.colab import drive
    if not os.path.exists('/content/drive/MyDrive'):
        drive.mount('/content/drive')
except Exception:
    pass
os.makedirs(OUT_DIR, exist_ok=True)
os.makedirs(WORK_DIR, exist_ok=True)

_SA_KEY = os.environ.get('GEE_SERVICE_ACCOUNT_KEY', '').strip()
if _SA_KEY:
    # Unattended (e.g. GitHub Actions): service-account JSON key, as text or a file path
    import json
    _key_text = _SA_KEY if _SA_KEY.startswith('{') else open(_SA_KEY, encoding='utf-8').read()
    _credentials = ee.ServiceAccountCredentials(json.loads(_key_text)['client_email'], key_data=_key_text)
    ee.Initialize(_credentials, project=PROJECT_ID)
    print(f"Earth Engine: service account {json.loads(_key_text)['client_email']}")
else:
    try:
        ee.Initialize(project=PROJECT_ID)
    except Exception:
        ee.Authenticate()
        ee.Initialize(project=PROJECT_ID)

RUN_STAMP = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')


def empty_raw():
    return pd.DataFrame(columns=RAW_COLS)


# =========================================================================== #
# DOWNLOAD FUNCTIONS - DAHITI (incremental per station)                       #
# =========================================================================== #
def dahiti_type(name, raw_type=''):
    text = (str(raw_type or '') + ' ' + str(name or '')).lower()
    return 'lake' if any(k in text for k in ('lake', 'reservoir', 'wetland', 'swamp')) else 'river'


def seed_from_old_combined_file():
    """First run only: reuse the newest combined_water_levels_analyzed*.xlsx made by the
    old DAHITI script (it is read, never deleted)."""
    old = sorted(glob.glob(os.path.join(OUT_DIR, 'combined_water_levels_analyzed*.xlsx')),
                 key=os.path.getmtime)
    if DAHITI_SEED_FILE:
        if os.path.exists(DAHITI_SEED_FILE):
            old = [DAHITI_SEED_FILE]
        else:
            print(f"DAHITI_SEED_FILE not found ({DAHITI_SEED_FILE}); looking for the newest old file instead.")
    if not old:
        return pd.DataFrame(columns=DAHITI_RAW_COLS)
    path = old[-1]
    try:
        d = pd.read_excel(path)
    except Exception as e:
        print(f"Could not read {os.path.basename(path)}: {e}")
        return pd.DataFrame(columns=DAHITI_RAW_COLS)
    d = d.rename(columns={'datetime': 'date [yyyy-mm-dd]', 'date': 'date [yyyy-mm-dd]',
                          'date_time': 'date [yyyy-mm-dd]', 'wse': 'water level'})
    d = d.loc[:, ~d.columns.duplicated()]
    need = ['DAHITI-ID', 'Target Name', 'Longitude', 'Latitude', 'date [yyyy-mm-dd]', 'water level']
    if not all(c in d.columns for c in need):
        print(f"{os.path.basename(path)} is missing expected columns; not used.")
        return pd.DataFrame(columns=DAHITI_RAW_COLS)
    if 'wse_u' not in d:
        d['wse_u'] = np.nan
    d['Target Type'] = d['Target Name'].map(dahiti_type)
    d['date [yyyy-mm-dd]'] = pd.to_datetime(d['date [yyyy-mm-dd]'], errors='coerce')
    d = d[DAHITI_RAW_COLS].dropna(subset=['date [yyyy-mm-dd]', 'water level'])
    d = d[d['DAHITI-ID'].isin(DAHITI_IDS)].drop_duplicates(['DAHITI-ID', 'date [yyyy-mm-dd]'])
    print(f"Seeded {len(d)} rows for {d['DAHITI-ID'].nunique()} stations from old file "
          f"{os.path.basename(path)}")
    return d


def download_dahiti():
    global DAHITI_RAW_XLSX
    print("\n--- DAHITI (only dates after the last saved one, per station) ---")
    if DAHITI_SEED_FILE and os.path.abspath(DAHITI_RAW_XLSX) == os.path.abspath(DAHITI_SEED_FILE):
        DAHITI_RAW_XLSX = os.path.join(OUT_DIR, 'dahiti_water_levels_raw.xlsx')
        print("DAHITI_RAW_XLSX pointed at the old analysed file; saving raw data to "
              f"{os.path.basename(DAHITI_RAW_XLSX)} instead so the old file isn't overwritten.")
    existing = pd.DataFrame(columns=DAHITI_RAW_COLS)
    if os.path.exists(DAHITI_RAW_XLSX):
        existing = pd.read_excel(DAHITI_RAW_XLSX)
        existing['date [yyyy-mm-dd]'] = pd.to_datetime(existing['date [yyyy-mm-dd]'], errors='coerce')
        existing = existing[existing['DAHITI-ID'].isin(DAHITI_IDS)]
        print(f"Loaded {len(existing)} saved rows for {existing['DAHITI-ID'].nunique()} stations "
              f"from {os.path.basename(DAHITI_RAW_XLSX)}")
    if existing.empty:
        existing = seed_from_old_combined_file()
        if existing.empty:
            print("No saved DAHITI data yet - full download.")
        else:
            existing.to_excel(DAHITI_RAW_XLSX, index=False)

    if not DOWNLOAD_DAHITI:
        print("DOWNLOAD_DAHITI = False -> using saved data only.")
        return existing
    if not DAHITI_API_KEY:
        print("No DAHITI_API_KEY secret found -> using saved data only.")
        return existing

    # Station metadata from the API, falling back to what is already saved
    meta = {}
    try:
        r = requests.post(URL_LIST, json=DAHITI_LIST_ARGS, timeout=120)
        if r.status_code == 403:
            print("DAHITI rejected the API key (HTTP 403). Server said: "
                  f"{r.text.strip()[:300]}\n"
                  "-> Check the key (no spaces/quotes) or get a new one from your DAHITI account. "
                  "Continuing with saved data only.")
            return existing
        r.raise_for_status()
        for t in r.json().get('data', []):
            if t.get('dahiti_id') in DAHITI_IDS:
                meta[t['dahiti_id']] = {
                    'name': t.get('target_name'),
                    'type': dahiti_type(t.get('target_name'), t.get('type') or t.get('target_type')),
                    'lon': float(t['longitude']), 'lat': float(t['latitude'])}
        print(f"Metadata from API for {len(meta)}/{len(DAHITI_IDS)} stations")
    except Exception as e:
        print(f"list-targets failed ({e}); using saved station metadata where available.")
    if not existing.empty:
        for sid, grp in existing.groupby('DAHITI-ID'):
            last = grp.iloc[-1]
            meta.setdefault(int(sid), {
                'name': last['Target Name'],
                'type': last.get('Target Type') or dahiti_type(last['Target Name']),
                'lon': float(last['Longitude']), 'lat': float(last['Latitude'])})

    last_by_id = (existing.groupby('DAHITI-ID')['date [yyyy-mm-dd]'].max().to_dict()
                  if not existing.empty else {})
    # Stations missing from the raw file but already merged: start after the merged file's last date
    for sid in DAHITI_IDS:
        merged_last = MASTER_LAST.get(f'dahiti:{sid}')
        if merged_last is not None and (sid not in last_by_id or pd.isna(last_by_id[sid])
                                        or merged_last > last_by_id[sid]):
            last_by_id[sid] = merged_last

    frames = []
    for sid in DAHITI_IDS:
        m = meta.get(sid)
        if m is None:
            print(f"  {sid}: no metadata, skipped")
            continue
        args = {'api_key': DAHITI_API_KEY, 'dahiti_id': sid, 'format': 'csv'}
        last = last_by_id.get(sid)
        if last is not None and pd.notna(last):
            args['start_date'] = (last + pd.Timedelta(days=1)).strftime('%Y-%m-%d')
        try:
            r = requests.post(URL_DOWNLOAD, json=args, timeout=120)
            if r.status_code != 200:
                print(f"  {sid}: HTTP {r.status_code}")
                continue
            d = pd.read_csv(StringIO(r.text), sep=';')
        except Exception as e:
            print(f"  {sid}: {e}")
            continue
        if d.empty or 'datetime' not in d:
            print(f"  {sid}: no new data" + (f" since {args['start_date']}" if 'start_date' in args else ''))
            continue
        frames.append(pd.DataFrame({
            'DAHITI-ID': sid, 'Target Name': m['name'], 'Target Type': m['type'],
            'Longitude': m['lon'], 'Latitude': m['lat'],
            'date [yyyy-mm-dd]': pd.to_datetime(d['datetime'], errors='coerce'),
            'water level': pd.to_numeric(d['wse'], errors='coerce'),
            'wse_u': pd.to_numeric(d['wse_u'], errors='coerce') if 'wse_u' in d else np.nan,
        }))
        print(f"  {sid}: +{len(d)} rows")
        time.sleep(0.5)

    new = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=DAHITI_RAW_COLS)
    out = (pd.concat([existing, new], ignore_index=True)
             .drop_duplicates(['DAHITI-ID', 'date [yyyy-mm-dd]'], keep='last')
             .sort_values(['DAHITI-ID', 'date [yyyy-mm-dd]'])
             .reset_index(drop=True))
    if not out.empty:
        out.to_excel(DAHITI_RAW_XLSX, index=False)
    print(f"DAHITI saved: {out['DAHITI-ID'].nunique()} stations, {len(out)} rows "
          f"({len(new)} new) -> {DAHITI_RAW_XLSX}")
    return out


def dahiti_to_raw(d):
    if d.empty:
        return empty_raw()
    types = d['Target Type'] if 'Target Type' in d else pd.Series(np.nan, index=d.index)
    types = types.fillna(d['Target Name'].map(dahiti_type))
    return pd.DataFrame({
        'source': 'DAHITI', 'station_id': d['DAHITI-ID'].astype(int), 'type': types,
        'location': d['Target Name'], 'latitude': d['Latitude'].astype(float),
        'longitude': d['Longitude'].astype(float), 'river_key': None,
        'date_dt': pd.to_datetime(d['date [yyyy-mm-dd]'], errors='coerce'),
        'wse': pd.to_numeric(d['water level'], errors='coerce'),
        'uncertainty': pd.to_numeric(d['wse_u'], errors='coerce'),
    })


# =========================================================================== #
# DOWNLOAD FUNCTIONS - HYDROWEB.NEXT (rivers + lakes)                         #
# =========================================================================== #
def download_hydroweb():
    print("\n--- Hydroweb.next (rivers + lakes) ---")
    if not DOWNLOAD_HYDROWEB:
        print("DOWNLOAD_HYDROWEB = False -> using the zips already on Drive.")
        return
    if not HYDROWEB_API_KEY:
        print("No HYDROWEB_API_KEY secret found -> using the zips already on Drive.")
        return
    try:
        import py_hydroweb
    except ImportError:
        subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'py-hydroweb'], check=False)
        import py_hydroweb

    client = py_hydroweb.Client(api_key=HYDROWEB_API_KEY)
    # Download to local disk first (Drive root isn't writable and some py_hydroweb
    # versions write relative to the current folder), then copy to Drive.
    dl_dir = '/content/hydroweb_download' if os.path.isdir('/content') else os.path.abspath('hydroweb_download')
    os.makedirs(dl_dir, exist_ok=True)
    old_cwd = os.getcwd()
    for collection, target in [('HYDROWEB_RIVERS_OPE', RIVERS_ZIP), ('HYDROWEB_LAKES_OPE', LAKES_ZIP)]:
        print(f"Requesting {collection} for bbox {HYDROWEB_BBOX} ...")
        local_zip = os.path.join(dl_dir, f"{collection}_{RUN_STAMP}.zip")
        try:
            os.chdir(dl_dir)
            basket = py_hydroweb.DownloadBasket(f"{collection.lower()}_{RUN_STAMP}")
            basket.add_collection(collection, bbox=HYDROWEB_BBOX)
            path = client.submit_and_download_zip(basket, zip_filename=local_zip, output_folder=dl_dir)
            if path is None:
                print("  Hydroweb could not prepare this download (see message above); "
                      "keeping previous zip if present.")
                continue
            path = path if os.path.isabs(path) else os.path.join(dl_dir, path)
            if not os.path.exists(path):
                path = local_zip
            shutil.copy(path, target)          # only replace the Drive copy after a good download
            print(f"  saved -> {target} ({os.path.getsize(target) / 1e6:.1f} MB)")
        except Exception as e:
            print(f"  {collection} download failed ({e}); keeping previous zip if present.")
        finally:
            os.chdir(old_cwd)


RIVER_COLS = [
    'date', 'time', 'water_surface_elevation', 'uncertainty', 'sep',
    'lon', 'lat', 'ellipsoidal_height', 'geoid_ondulation', 'distance_km',
    'satellite', 'mission', 'ground_track', 'cycle', 'retracking', 'gdr_version',
]


def extract(zip_path, sub):
    out = os.path.join(WORK_DIR, sub)
    if not os.path.exists(zip_path):
        print(f"  {os.path.basename(zip_path)} not found, skipping {sub}")
        return []
    shutil.rmtree(out, ignore_errors=True)       # no stale files from earlier runs
    with zipfile.ZipFile(zip_path) as z:
        z.extractall(out)
    files = sorted(glob.glob(os.path.join(out, '**', '*.txt'), recursive=True))
    print(f"  {sub}: {len(files)} station files")
    return files


def load_hydroweb_rivers():
    frames = []
    for path in extract(RIVERS_ZIP, 'rivers'):
        m = {}
        with open(path, encoding='utf-8', errors='replace') as f:
            for line in f:
                if not line.startswith('#'):
                    break
                if '::' in line:
                    k, v = line[1:].split('::', 1)
                    m[k.strip()] = v.strip()
        if 'ID' not in m:
            continue
        d = pd.read_csv(path, comment='#', sep=r'\s+', header=None,
                        names=RIVER_COLS, dtype={'date': str, 'time': str})
        river = (m.get('RIVER') or '').replace('-', ' ').title()
        km = m.get('REFERENCE DISTANCE (km)')
        frames.append(pd.DataFrame({
            'source': 'Hydroweb', 'station_id': int(m['ID']), 'type': 'river',
            'location': f"{river} km {km}" if km else river,
            'latitude': float(m['REFERENCE LATITUDE']),
            'longitude': float(m['REFERENCE LONGITUDE']),
            'river_key': f"{m.get('BASIN', '')}|{m.get('RIVER', '')}".upper().replace('-', ' '),
            'date_dt': pd.to_datetime(d['date'] + ' ' + d['time'], errors='coerce'),
            'wse': pd.to_numeric(d['water_surface_elevation'], errors='coerce'),
            'uncertainty': pd.to_numeric(d['uncertainty'], errors='coerce'),
        }))
    return pd.concat(frames, ignore_index=True) if frames else empty_raw()


def load_hydroweb_lakes():
    frames = []
    for path in extract(LAKES_ZIP, 'lakes'):
        with open(path, encoding='utf-8', errors='replace') as f:
            lines = f.readlines()
        if not lines:
            continue
        h = {}
        for part in lines[0].strip().split(';'):
            if '=' in part:
                k, v = part.split('=', 1)
                h[k.strip()] = v.strip()
        if 'id' not in h:
            continue
        recs = []
        for l in lines[1:]:
            if not l.strip() or l.strip().startswith('#'):
                continue
            parts = [p.strip() for p in l.split(';')]
            if len(parts) >= 7:
                recs.append((parts + [''])[:8])
        if not recs:
            continue
        d = pd.DataFrame(recs, columns=['decimal_year', 'date', 'time', 'height_m',
                                        'stdev_m', 'area_km2', 'volume_km3', 'mission'])
        d['date'] = d['date'].str.replace('/', '-')
        name = h.get('lake', '').replace('-', ' ').title()
        frames.append(pd.DataFrame({
            'source': 'Hydroweb', 'station_id': int(h['id']), 'type': 'lake',
            'location': f"{name} Lake",
            'latitude': float(h.get('lat', 'nan')), 'longitude': float(h.get('lon', 'nan')),
            'river_key': None,
            'date_dt': pd.to_datetime(d['date'] + ' ' + d['time'], errors='coerce'),
            'wse': pd.to_numeric(d['height_m'], errors='coerce'),
            'uncertainty': pd.to_numeric(d['stdev_m'], errors='coerce'),
        }))
    return pd.concat(frames, ignore_index=True) if frames else empty_raw()


def load_hydroweb():
    rivers, lakes = load_hydroweb_rivers(), load_hydroweb_lakes()
    hw = pd.concat([rivers, lakes], ignore_index=True)
    if hw.empty:
        print("Hydroweb: no data.")
        return hw
    mn_lon, mn_lat, mx_lon, mx_lat = HYDROWEB_BBOX
    hw = hw[hw.longitude.between(mn_lon, mx_lon) & hw.latitude.between(mn_lat, mx_lat)]

    # Save the parsed raw download, like the previous Hydroweb script did
    with pd.ExcelWriter(HYDROWEB_RAW_XLSX, engine='openpyxl') as xw:
        stations = (hw.groupby(['type', 'station_id'])
                      .agg(location=('location', 'first'), latitude=('latitude', 'first'),
                           longitude=('longitude', 'first'), n_measurements=('wse', 'size'),
                           first_date=('date_dt', 'min'), last_date=('date_dt', 'max'))
                      .reset_index())
        stations.to_excel(xw, sheet_name='Stations', index=False)
        for t, sheet in [('river', 'All_Rivers'), ('lake', 'All_Lakes')]:
            hw[hw.type == t].drop(columns=['river_key']).to_excel(xw, sheet_name=sheet, index=False)
    print(f"Hydroweb saved: {(hw.type == 'river').groupby(hw.station_id).any().sum()} river + "
          f"{(hw.type == 'lake').groupby(hw.station_id).any().sum()} lake stations, "
          f"{len(hw)} rows -> {HYDROWEB_RAW_XLSX}")
    return hw


# =========================================================================== #
# STEP 1 - CHECK MERGED STATION OBSERVATIONS                                  #
# =========================================================================== #
WSE_COL = 'Water Surface Elevation - values(m)'


def obs_key(frame):
    """One key per observation (output-table columns): source | station | date | level (mm)."""
    return (frame['source'].astype(str) + '|'
            + pd.to_numeric(frame['station_id']).astype(int).astype(str) + '|'
            + frame['date'].astype(str) + '|'
            + pd.to_numeric(frame[WSE_COL]).round(3).map('{:.3f}'.format))


print("\n=== STEP 1: CHECK MERGED STATION OBSERVATIONS ===")
if REBUILD_OUTPUTS:
    for p in (MASTER_CSV, MASTER_XLSX, GEE_LEDGER):
        if os.path.exists(p):
            os.remove(p)
    print("REBUILD_OUTPUTS = True -> master files and GEE ledger removed; rebuilding all rows.")

MASTER_LAST = {}          # station_uid -> last saved observation date
master_hist = None
if os.path.exists(MASTER_CSV):
    master_hist = pd.read_csv(MASTER_CSV, usecols=['source', 'station_id', 'type', 'location/river_name',
                                                   'latitude', 'longitude', 'date', WSE_COL,
                                                   'uncertainty (m)'])
    old_keys = set(obs_key(master_hist))
    master_hist['station_uid'] = (master_hist['source'].str.lower() + ':'
                                  + master_hist['station_id'].astype(int).astype(str))
    MASTER_LAST = pd.to_datetime(master_hist['date']).groupby(master_hist['station_uid']).max().to_dict()
    summary = master_hist.groupby('source').agg(stations=('station_uid', 'nunique'),
                                                observations=('date', 'size'),
                                                first_date=('date', 'min'), last_date=('date', 'max'))
    print(f"{os.path.basename(MASTER_CSV)}: {len(master_hist)} observations already merged")
    print(summary.to_string())
else:
    old_keys = set()
    print("No merged file yet -> every downloaded observation counts as new.")

# =========================================================================== #
# STEP 2 - DOWNLOAD NEW OBSERVATIONS FROM THE DAHITI + HYDROWEB APIs           #
# =========================================================================== #
print("\n=== STEP 2: DOWNLOAD NEW OBSERVATIONS (DAHITI + Hydroweb APIs) ===")
dahiti_raw = download_dahiti()
download_hydroweb()
hydroweb_raw = load_hydroweb()

# =========================================================================== #
# STEP 3 - MERGE, THEN GEE CLIMATE + ELEVATION FOR THE NEW OBSERVATIONS        #
# =========================================================================== #
print("\n=== STEP 3: MERGE + GEE CLIMATE FOR NEW OBSERVATIONS ===")
raw = pd.concat([dahiti_to_raw(dahiti_raw), hydroweb_raw], ignore_index=True)

# The merged file is also history: running stats and lags for new rows need all
# earlier observations, even if a raw download file has gone missing.
if master_hist is not None and len(master_hist):
    hist = pd.DataFrame({
        'source': master_hist['source'], 'station_id': master_hist['station_id'].astype(int),
        'type': master_hist['type'], 'location': master_hist['location/river_name'],
        'latitude': master_hist['latitude'], 'longitude': master_hist['longitude'], 'river_key': None,
        'date_dt': pd.to_datetime(master_hist['date']), 'wse': master_hist[WSE_COL],
        'uncertainty': master_hist['uncertainty (m)']})
    raw = pd.concat([raw, hist], ignore_index=True)          # downloaded rows first -> kept on duplicates

raw['date_dt'] = pd.to_datetime(raw['date_dt'], errors='coerce')
for col in ['wse', 'uncertainty', 'latitude', 'longitude']:       # an empty source can leave these as text
    raw[col] = pd.to_numeric(raw[col], errors='coerce')
raw = raw.dropna(subset=['date_dt', 'wse', 'latitude', 'longitude'])
raw = raw[raw.date_dt.dt.year >= START_YEAR].copy()
raw['station_id'] = raw['station_id'].astype(int)
raw['station_uid'] = raw['source'].str.lower() + ':' + raw['station_id'].astype(str)
raw['river_key'] = raw.groupby('station_uid')['river_key'].transform('first')
raw['_k'] = obs_key(pd.DataFrame({'source': raw['source'], 'station_id': raw['station_id'],
                                  'date': raw['date_dt'].dt.strftime('%Y-%m-%d'), WSE_COL: raw['wse']}))
raw = raw.drop_duplicates('_k').drop_duplicates(['station_uid', 'date_dt']).drop(columns='_k')
df = raw.sort_values(['station_uid', 'date_dt']).reset_index(drop=True)

if df.empty:
    raise SystemExit("No data from either source - nothing to merge.")

df['day_dt'] = df['date_dt'].dt.normalize()
df['date_key'] = df['day_dt'].dt.strftime('%Y-%m-%d')
df['obs_key'] = obs_key(pd.DataFrame({'source': df['source'], 'station_id': df['station_id'],
                                      'date': df['date_key'], WSE_COL: df['wse']}))
is_new = ~df['obs_key'].isin(old_keys)
new_by_source = df.loc[is_new].groupby('source')['obs_key'].size().to_dict()
print(f"New observations not yet merged: {int(is_new.sum())} {new_by_source if new_by_source else ''}")
if not is_new.any():
    print("Nothing new -> no GEE climate requests; Step 4 will only retry any pending uploads.")

# --- 3a. Station statistics, "as of" each observation -------------------------
# Running values from the station's first observation up to and including this
# one, so a saved row never changes when newer data arrive.
g = df.groupby('station_uid')['wse']
df['minimium'] = g.cummin()
df['maxmium'] = g.cummax()
df['average'] = g.cumsum() / (df.groupby('station_uid').cumcount() + 1)
df['deviation from average'] = df['wse'] - df['average']
df['difference from previous'] = g.diff().fillna(0.0)

stations = (df.groupby('station_uid')
              .agg(source=('source', 'first'), station_id=('station_id', 'first'),
                   type=('type', 'first'), location=('location', 'first'),
                   latitude=('latitude', 'first'), longitude=('longitude', 'first'),
                   river_key=('river_key', 'first'), mean_wse=('wse', 'mean'),
                   n_obs=('wse', 'size'), first_date=('date_key', 'min'),
                   last_date=('date_key', 'max'))
              .reset_index())
present = set(stations.station_uid)
print(f"Merged: {len(stations)} stations, {len(df)} observations")

# --- 3b. Upstream / downstream links -----------------------------------------
raw_links = {}
order = pd.read_csv(StringIO(DAHITI_ORDER), sep='\t').dropna()
for _, r in order.iterrows():
    raw_links[f"dahiti:{int(r['id'])}"] = f"dahiti:{int(r['after'])}"

# Hydroweb rivers: chain stations on the same river, highest mean level = upstream
hw_riv = stations[(stations.source == 'Hydroweb') & (stations.type == 'river')]
for _, grp in hw_riv.groupby('river_key'):
    uids = grp.sort_values('mean_wse', ascending=False).station_uid.tolist()
    for up, dn in zip(uids[:-1], uids[1:]):
        raw_links.setdefault(up, dn)

for up, dn in MANUAL_LINKS:
    raw_links[up] = dn


def resolve_down(uid):
    """Follow the chain past stations that have no data (e.g. 38471, 38454)."""
    seen, nxt = {uid}, raw_links.get(uid)
    while nxt is not None and nxt not in present and nxt not in seen:
        seen.add(nxt)
        nxt = raw_links.get(nxt)
    return nxt if nxt in present else None


after_map = {u: resolve_down(u) for u in present}
after_map = {u: d for u, d in after_map.items() if d is not None}
before_map = {}
for up, dn in after_map.items():
    before_map.setdefault(dn, []).append(up)

stations['Station After_this_station'] = stations.station_uid.map(after_map)
stations['Station Before_this_Station'] = stations.station_uid.map(
    lambda u: ', '.join(sorted(before_map[u])) if u in before_map else None)

# --- 3c. Onset lag over the last 12 observations -------------------------------
# For an observation on day t at station A (downstream station B):
#   window = from A's 12th-last observation up to t (both included)
#   Lag Days = the shift L (0..MAX_LAG_DAYS days) at which B's daily water level
#              inside the window best matches A's level L days earlier.
# Only data up to day t are used, so the value is fixed once written.
print("Computing onset lags (rolling window of the last "
      f"{LAG_WINDOW_OBS} observations)...")
_daily_cache = {}


def daily_change(uid):
    """Daily water level, linearly interpolated between passes (not across gaps > MAX_GAP_DAYS).
    Levels rather than day-to-day changes: with 10-35-day revisits the changes are step
    functions and lock onto pass timing; levels recover the true delay far better."""
    if uid in _daily_cache:
        return _daily_cache[uid]
    s = df.loc[df.station_uid == uid].groupby('day_dt')['wse'].mean()
    if len(s) < 3:
        _daily_cache[uid] = None
        return None
    idx = pd.date_range(s.index.min(), s.index.max(), freq='D')
    obs = pd.Series(s.index, index=s.index).reindex(idx)
    gap = (obs.bfill() - obs.ffill()).dt.days
    full = s.reindex(idx).interpolate(method='time', limit_area='inside')
    full[gap > MAX_GAP_DAYS] = np.nan
    _daily_cache[uid] = full
    return full


def best_lag(up_vals, dn_vals, s, e):
    """Best L for downstream[s..e] vs upstream[s-L..e-L] (positions on a shared daily grid)."""
    y = dn_vals[s:e + 1]
    best_L, best_r = np.nan, -np.inf
    for L in range(MAX_LAG_DAYS + 1):
        x = up_vals[s - L:e + 1 - L]
        m = np.isfinite(x) & np.isfinite(y)
        if m.sum() < LAG_MIN_OVERLAP_DAYS:
            continue
        xs, ys = x[m] - x[m].mean(), y[m] - y[m].mean()
        if xs.std() == 0 or ys.std() == 0:
            continue
        r = np.corrcoef(xs, ys)[0, 1]
        if np.isfinite(r) and r > best_r:
            best_L, best_r = L, r
    if not np.isfinite(best_r):
        return np.nan, np.nan
    return (best_L if best_r >= MIN_R else np.nan), best_r


# Rows that need a lag: those not saved yet (saved rows keep their value)
need_lag = is_new
df['Lag Days'] = np.nan
df['lag_r'] = np.nan
for up, grp in df[need_lag].groupby('station_uid'):
    dn = after_map.get(up)
    su, sd = daily_change(up), (daily_change(dn) if dn else None)
    if su is None or sd is None:
        continue
    start = min(su.index.min(), sd.index.min()) - pd.Timedelta(days=MAX_LAG_DAYS)
    grid = pd.date_range(start, max(su.index.max(), sd.index.max()), freq='D')
    up_vals, dn_vals = su.reindex(grid).to_numpy(), sd.reindex(grid).to_numpy()
    pos = pd.Series(np.arange(len(grid)), index=grid)
    obs_days = np.sort(df.loc[df.station_uid == up, 'day_dt'].unique())
    for i, day in grp['day_dt'].items():
        k = np.searchsorted(obs_days, np.datetime64(day), side='right')   # obs up to and incl. this day
        if k < LAG_WINDOW_OBS:
            continue
        s, e = pos[pd.Timestamp(obs_days[k - LAG_WINDOW_OBS])], pos[day]
        df.at[i, 'Lag Days'], df.at[i, 'lag_r'] = best_lag(up_vals, dn_vals, s, e)

# Stations layer keeps its original 'median_lag_days' column: median over saved + new rows
all_lags = df.loc[need_lag, ['station_uid', 'Lag Days']]
if os.path.exists(MASTER_CSV):
    saved = pd.read_csv(MASTER_CSV, usecols=['source', 'station_id', 'Lag Days'])
    saved['station_uid'] = saved['source'].str.lower() + ':' + saved['station_id'].astype(int).astype(str)
    all_lags = pd.concat([saved[['station_uid', 'Lag Days']], all_lags], ignore_index=True)
stations['median_lag_days'] = stations.station_uid.map(all_lags.groupby('station_uid')['Lag Days'].median())
print(f"  Lag Days computed for {int(need_lag.sum())} new rows; "
      f"{int(df.loc[need_lag, 'Lag Days'].notna().sum())} with an accepted lag")

# --- 3d. Elevation (SRTM, once per station, cached) --------------------------
elev_cols = ['Elevation_Min', 'Elevation_Mean', 'Elevation_Max']
elev = (pd.read_csv(ELEV_CACHE) if os.path.exists(ELEV_CACHE)
        else pd.DataFrame(columns=['station_uid'] + elev_cols))
todo = stations[~stations.station_uid.isin(elev.station_uid)]
if len(todo):
    print(f"Fetching SRTM elevation for {len(todo)} stations...")
    fc = ee.FeatureCollection([
        ee.Feature(ee.Geometry.Point([float(r.longitude), float(r.latitude)]).buffer(ELEV_BUFFER_M),
                   {'uid': r.station_uid}) for r in todo.itertuples()])
    red = (ee.Reducer.min().combine(ee.Reducer.mean(), '', True)
           .combine(ee.Reducer.max(), '', True))
    res = (ee.Image('USGS/SRTMGL1_003').select('elevation')
           .reduceRegions(collection=fc, reducer=red, scale=30).getInfo())
    new = pd.DataFrame([f['properties'] for f in res['features']]).rename(
        columns={'uid': 'station_uid', 'min': 'Elevation_Min',
                 'mean': 'Elevation_Mean', 'max': 'Elevation_Max'})
    elev = pd.concat([elev, new[['station_uid'] + elev_cols]], ignore_index=True)
    elev.to_csv(ELEV_CACHE, index=False)
stations = stations.merge(elev, on='station_uid', how='left')

# --- 3e. 15-day antecedent hydro-climate (CHIRPS + ERA5-Land, cached) --------
CHIRPS = ee.ImageCollection('UCSB-CHG/CHIRPS/DAILY').select(['precipitation'], ['rain_sum'])
ERA5 = ee.ImageCollection('ECMWF/ERA5_LAND/DAILY_AGGR')


def climate_window(f):
    end = ee.Date(f.get('date')).advance(1, 'day')          # include the observation day
    start = end.advance(-WINDOW_DAYS, 'day')
    roi = f.geometry().buffer(ee.Number(f.get('buf')))
    ch = CHIRPS.filterDate(start, end)
    er = ERA5.filterDate(start, end)
    img = ee.Image.cat([
        ch.sum(),
        er.select(['temperature_2m', 'volumetric_soil_water_layer_1'],
                  ['t2m_mean', 'swvl1_mean']).mean(),
        er.select(['total_evaporation_sum', 'runoff_sum'],
                  ['evap_sum', 'ro_sum']).sum(),
    ])
    stats = img.reduceRegion(reducer=ee.Reducer.mean(), geometry=roi,
                             scale=5000, maxPixels=1e9, tileScale=2)
    return (ee.Feature(None, {'uid': f.get('uid'), 'date': f.get('date')})
            .setMulti(stats)
            .set({'n_chirps': ch.size(), 'n_era5': er.size()}))


clim_cols = ['station_uid', 'date_key', 'rain_sum', 't2m_mean', 'swvl1_mean',
             'evap_sum', 'ro_sum', 'n_chirps', 'n_era5', 'complete']
clim = (pd.read_csv(CLIMATE_CACHE) if os.path.exists(CLIMATE_CACHE)
        else pd.DataFrame(columns=clim_cols))

need = (df.loc[is_new, ['station_uid', 'type', 'latitude', 'longitude', 'date_key']]   # new observations only
        .drop_duplicates(['station_uid', 'date_key'])
        .merge(clim[['station_uid', 'date_key', 'complete']], on=['station_uid', 'date_key'], how='left'))
cutoff = pd.Timestamp.today().normalize() - pd.Timedelta(days=RETRY_RECENT_DAYS)
is_recent = pd.to_datetime(need.date_key) >= cutoff
todo = need[need.complete.isna() | ((need.complete.astype(str) == 'False') & is_recent)]

if len(todo):
    n_batches = (len(todo) + GEE_BATCH - 1) // GEE_BATCH
    print(f"Fetching 15-day climate for {len(todo)} station-dates in {n_batches} batches...")
    for b in range(n_batches):
        batch = todo.iloc[b * GEE_BATCH:(b + 1) * GEE_BATCH]
        fc = ee.FeatureCollection([
            ee.Feature(ee.Geometry.Point([float(r.longitude), float(r.latitude)]),
                       {'uid': r.station_uid, 'date': r.date_key,
                        'buf': BUFFER_M.get(r.type, 5000)})
            for r in batch.itertuples()])
        result = None
        for attempt in range(3):
            try:
                result = fc.map(climate_window).getInfo()['features']
                break
            except Exception as e:
                print(f"  batch {b + 1}/{n_batches} attempt {attempt + 1} failed: {e}")
                time.sleep(10)
        if result is None:
            continue
        new = pd.DataFrame([f['properties'] for f in result]).rename(
            columns={'uid': 'station_uid', 'date': 'date_key'})
        for c in clim_cols:
            if c not in new:
                new[c] = np.nan
        vals = ['rain_sum', 't2m_mean', 'swvl1_mean', 'evap_sum', 'ro_sum']
        new['complete'] = ((new.n_chirps >= WINDOW_DAYS) & (new.n_era5 >= WINDOW_DAYS)
                           & new[vals].notna().all(axis=1))
        keys = set(zip(new.station_uid, new.date_key))
        keep = [k not in keys for k in zip(clim.station_uid, clim.date_key)]
        clim = pd.concat([clim[keep], new[clim_cols]], ignore_index=True)
        clim.to_csv(CLIMATE_CACHE, index=False)            # save progress every batch
        print(f"  batch {b + 1}/{n_batches} done")

c = clim.copy()
for col in ['rain_sum', 't2m_mean', 'swvl1_mean', 'evap_sum', 'ro_sum', 'n_chirps', 'n_era5']:
    c[col] = pd.to_numeric(c[col], errors='coerce')
full_ch = c.n_chirps >= WINDOW_DAYS
full_er = c.n_era5 >= WINDOW_DAYS
c['last_15_days_Rainfall_mm'] = c.rain_sum.where(full_ch)
c['last_15_days_Soil_Moisture'] = c.swvl1_mean.where(full_er)             # m3/m3
c['last_15_days_Evapotranspiration'] = (-c.evap_sum * 1000).where(full_er)  # mm, ERA5 sign flipped
c['last_15_days_Temp_C'] = (c.t2m_mean - 273.15).where(full_er)
c['last_15_days_runoff'] = (c.ro_sum * 1000).where(full_er)                # mm
clim_out = c[['station_uid', 'date_key', 'last_15_days_Rainfall_mm', 'last_15_days_Soil_Moisture',
              'last_15_days_Evapotranspiration', 'last_15_days_Temp_C', 'last_15_days_runoff']]
df = df.merge(clim_out, on=['station_uid', 'date_key'], how='left')

# --- 3f. Final table + export ------------------------------------------------
df = df.merge(stations[['station_uid'] + elev_cols
                       + ['Station After_this_station', 'Station Before_this_Station']],
              on='station_uid', how='left')

final = pd.DataFrame({
    'source': df['source'],
    'station_id': df['station_id'],
    'type': df['type'],
    'location/river_name': df['location'],
    'latitude': df['latitude'],
    'longitude': df['longitude'],
    'year': df['date_dt'].dt.year,
    'month': df['date_dt'].dt.month,
    'day': df['date_dt'].dt.day,
    'date': df['date_key'],
    'Water Surface Elevation - values(m)': df['wse'],
    'uncertainty (m)': df['uncertainty'],
    'minimium': df['minimium'],
    'maxmium': df['maxmium'],
    'average': df['average'],
    'deviation from average': df['deviation from average'],
    'difference from previous': df['difference from previous'],
    'last_15_days_Rainfall_mm': df['last_15_days_Rainfall_mm'],
    'last_15_days_Soil_Moisture': df['last_15_days_Soil_Moisture'],
    'last_15_days_Evapotranspiration': df['last_15_days_Evapotranspiration'],
    'last_15_days_Temp_C': df['last_15_days_Temp_C'],
    'last_15_days_runoff': df['last_15_days_runoff'],
    'Elevation_Max': df['Elevation_Max'],
    'Elevation_Mean': df['Elevation_Mean'],
    'Elevation_Min': df['Elevation_Min'],
    'Station After_this_station': df['Station After_this_station'],
    'Station Before_this_Station': df['Station Before_this_Station'],
    'Lag Days': df['Lag Days'],
}).sort_values(['source', 'type', 'station_id', 'date']).reset_index(drop=True)

notes = pd.DataFrame([
    ('station links', "IDs written as source:id. DAHITI links from the manual order table; "
                      "Hydroweb river stations chained along the same river by mean level "
                      "(higher = upstream); MANUAL_LINKS override both."),
    ('last_15_days_*', f"{WINDOW_DAYS} days ending on and including the observation date; "
                       f"buffer {BUFFER_M['river']} m (river) / {BUFFER_M['lake']} m (lake). "
                       "Blank if the window isn't fully covered yet."),
    ('last_15_days_Rainfall_mm', 'CHIRPS daily, summed (mm)'),
    ('last_15_days_Soil_Moisture', 'ERA5-Land volumetric_soil_water_layer_1, mean (m3/m3, 0-7 cm)'),
    ('last_15_days_Evapotranspiration', 'ERA5-Land total_evaporation_sum, summed, sign flipped (mm)'),
    ('last_15_days_Temp_C', 'ERA5-Land temperature_2m, mean (deg C)'),
    ('last_15_days_runoff', 'ERA5-Land runoff_sum, summed (mm)'),
    ('Elevation_*', f'SRTM 30 m within {ELEV_BUFFER_M} m of the station (m)'),
    ('minimium / maxmium / average', "Running values from the station's first observation "
                                     "up to and including this one (as of this date)."),
    ('deviation from average', 'This level minus the running average as of this observation (m)'),
    ('Lag Days', f"Days (0-{MAX_LAG_DAYS}) by which the downstream station's daily water level "
                 f"best follows this station's, over the window covered by this station's last "
                 f"{LAG_WINDOW_OBS} observations (data up to this date only). Blank if fewer than "
                 f"{LAG_WINDOW_OBS} observations, no downstream station, too little overlap or "
                 f"correlation below {MIN_R}."),
    ('difference from previous', 'Change from the previous observation at the same station (m)'),
], columns=['column', 'definition'])

CLIMATE_COLS_OUT = ['last_15_days_Rainfall_mm', 'last_15_days_Soil_Moisture',
                    'last_15_days_Evapotranspiration', 'last_15_days_Temp_C', 'last_15_days_runoff']


def to_py(v):
    return v.item() if hasattr(v, 'item') else v


# Rows whose 15-day climate window isn't complete yet (ERA5 lags ~1 week) are held
# back while they are recent, so they are written once, already filled in.
recent_row = pd.to_datetime(final['date']) >= cutoff
incomplete_row = final[CLIMATE_COLS_OUT].isna().any(axis=1)
ready = final[~(recent_row & incomplete_row)]
held_back = int((recent_row & incomplete_row).sum())

master_cols = (pd.read_csv(MASTER_CSV, nrows=0).columns.tolist() if os.path.exists(MASTER_CSV)
               else list(final.columns))
new_rows = ready[~obs_key(ready).isin(old_keys)].reindex(columns=master_cols)

# --- CSV: append ---
if not os.path.exists(MASTER_CSV):
    new_rows.to_csv(MASTER_CSV, index=False)
elif len(new_rows):
    new_rows.to_csv(MASTER_CSV, mode='a', header=False, index=False)

# --- XLSX: append to 'Merged', refresh the small summary sheets ---
if not os.path.exists(MASTER_XLSX):
    with pd.ExcelWriter(MASTER_XLSX, engine='openpyxl') as xw:
        new_rows.to_excel(xw, sheet_name='Merged', index=False)
        stations.drop(columns=['river_key']).to_excel(xw, sheet_name='Stations', index=False)
        notes.to_excel(xw, sheet_name='Column_notes', index=False)
else:
    from openpyxl import load_workbook
    wb = load_workbook(MASTER_XLSX)
    ws = wb['Merged']
    header = [c.value for c in ws[1]]
    for row in new_rows.reindex(columns=header).itertuples(index=False):
        ws.append([None if pd.isna(v) else to_py(v) for v in row])
    for name in ['Stations', 'Lags_by_year', 'Column_notes']:
        if name in wb.sheetnames:
            del wb[name]
    wb.save(MASTER_XLSX)
    with pd.ExcelWriter(MASTER_XLSX, engine='openpyxl', mode='a') as xw:
        stations.drop(columns=['river_key']).to_excel(xw, sheet_name='Stations', index=False)
        notes.to_excel(xw, sheet_name='Column_notes', index=False)

print("\nDone.")
print(final.groupby(['source', 'type'])['station_id'].nunique().rename('stations'))
print(f"New observations appended: {len(new_rows)}  "
      f"(already saved: {len(old_keys)}, held back until climate is complete: {held_back})")
print(f"-> {MASTER_XLSX}\n-> {MASTER_CSV}")


# =========================================================================== #
# STEP 4 - UPLOAD NEW STATION OBSERVATIONS TO EARTH ENGINE                    #
# =========================================================================== #
# Creates, under ASSET_FOLDER:
#   stations             one point per station (metadata, links, elevation, median lag)
#   merged_observations  one point per observation, every merged column, plus
#                        system:time_start so .filterDate() works in GEE
# Column names are made GEE-safe: "Water Surface Elevation - values(m)" ->
# "Water_Surface_Elevation_values_m", "Lag Days" -> "Lag_Days", etc.
def gee_name(col):
    return re.sub(r'[^0-9A-Za-z_]+', '_', str(col)).strip('_')


def df_to_fc(frame, date_col=None):
    feats = []
    for rec in frame.to_dict('records'):
        props = {gee_name(k): to_py(v) for k, v in rec.items() if pd.notna(v)}
        if date_col and pd.notna(rec.get(date_col)):
            props['system:time_start'] = int(pd.Timestamp(rec[date_col]).value // 10**6)
        geom = ee.Geometry.Point([float(rec['longitude']), float(rec['latitude'])])
        feats.append(ee.Feature(geom, props))
    return ee.FeatureCollection(feats)


def asset_exists(asset_id):
    try:
        ee.data.getAsset(asset_id)
        return True
    except ee.EEException:
        return False


def start_export(fc, asset_id, desc):
    if asset_exists(asset_id):
        ee.data.deleteAsset(asset_id)            # exports can't overwrite in place
    task = ee.batch.Export.table.toAsset(collection=fc, description=gee_name(desc)[:100],
                                         assetId=asset_id)
    task.start()
    return task


def wait_for(tasks, poll=30):
    done = ('COMPLETED', 'FAILED', 'CANCELLED')
    while True:
        states = [t.status()['state'] for t in tasks]
        if all(s in done for s in states):
            for t, s in zip(tasks, states):
                if s != 'COMPLETED':
                    print(f"  {t.status().get('description')}: {s} "
                          f"{t.status().get('error_message', '')}")
            return states
        print(f"  waiting: {sum(s == 'COMPLETED' for s in states)}/{len(tasks)} tasks done")
        time.sleep(poll)


if UPLOAD_TO_GEE:
    print("\n=== STEP 4: UPLOAD NEW STATION OBSERVATIONS TO GEE ===")
    if not asset_exists(ASSET_FOLDER):
        ee.data.createAsset({'type': 'FOLDER'}, ASSET_FOLDER)
        print(f"Created folder {ASSET_FOLDER}")
    target = f'{ASSET_FOLDER}/merged_observations'

    # Stations: small, so the whole layer is refreshed each run
    st_task = start_export(df_to_fc(stations.drop(columns=['river_key'])),
                           f'{ASSET_FOLDER}/stations', 'altimetry_stations')
    print(f"Started stations export ({len(stations)} points)")

    # Which saved observations are not in GEE yet? (the ledger is only updated after
    # a successful upload, so rows from a failed run are retried next time)
    master = pd.read_csv(MASTER_CSV)
    master_keys = obs_key(master)
    if os.path.exists(GEE_LEDGER):
        done_keys = set(pd.read_csv(GEE_LEDGER)['key'])
        rebuild = not asset_exists(target)
    else:
        done_keys, rebuild = set(), True   # first incremental run: build the asset once from the master file
    pending = master[~master_keys.isin(done_keys)] if not rebuild else master
    print(f"Observations to upload: {len(pending)}"
          + (" (building the asset from the full master file)" if rebuild else ""))

    if pending.empty:
        wait_for([st_task])
        print("Nothing new for GEE; stations layer refreshed.")
    else:
        # Same part names as the version that uploaded successfully: merged_obs_part_000, _001, ...
        n_parts = math.ceil(len(pending) / ROWS_PER_PART)
        part_ids, part_tasks = [], []
        for i in range(n_parts):
            chunk = pending.iloc[i * ROWS_PER_PART:(i + 1) * ROWS_PER_PART]
            pid = f'{ASSET_FOLDER}/merged_obs_part_{i:03d}'
            part_tasks.append(start_export(df_to_fc(chunk, date_col='date'), pid,
                                           f'merged_obs_part_{i:03d}'))
            part_ids.append(pid)
        print(f"Started {n_parts} observation export(s) of up to {ROWS_PER_PART} rows")
        states = wait_for(part_tasks + [st_task])

        if all(s == 'COMPLETED' for s in states[:-1]):
            new_fc = ee.FeatureCollection([ee.FeatureCollection(p) for p in part_ids]).flatten()
            # Append on Google's side: existing merged_observations + new rows, written to a
            # temporary copy first so the existing asset is only replaced after success
            combined = new_fc if rebuild else ee.FeatureCollection(target).merge(new_fc)
            staging = f'{target}_{RUN_STAMP}'
            t = start_export(combined, staging, 'merged_observations')
            if wait_for([t])[0] == 'COMPLETED':
                if asset_exists(target):
                    ee.data.deleteAsset(target)
                ee.data.renameAsset(staging, target)
                for p in part_ids:
                    ee.data.deleteAsset(p)
                uploaded = set(master_keys) if rebuild else done_keys | set(obs_key(pending))
                pd.DataFrame({'key': sorted(uploaded)}).to_csv(GEE_LEDGER, index=False)
                print(f"Done -> {target} (+{len(pending)} observations, {len(uploaded)} in total) "
                      f"and {ASSET_FOLDER}/stations")
            else:
                print("Merge step failed; the existing asset is unchanged and the new rows "
                      "will be retried next run.")
        else:
            print("Some uploads failed; the existing asset is unchanged and the new rows "
                  "will be retried next run.")
            for p in part_ids:
                if asset_exists(p):
                    ee.data.deleteAsset(p)
