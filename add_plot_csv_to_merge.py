# -*- coding: utf-8 -*-
"""Add Hydroweb 'plot_L_<lake>.csv' downloads (the CSV you get from the Hydroweb.next plot page) to
merged_altimetry_stations.csv and .xlsx. Run by hand: Actions -> "Add plot CSV lakes to merge (manual)".

What it writes: only the observation columns (source, station_id, type, location, latitude, longitude, date,
level, uncertainty, station_uid). The next nightly run reads these saved rows back as history, so it fills in
the derived columns (min / max / average, seasonal baseline, QC), the 15-day climate, elevation and the
GEE asset, exactly as for any other row.

Rules
  * one row per lake per day (the lowest-uncertainty pass of that day)
  * a day that is already in the merged file for that station is left alone (existing rows win)
  * rows before START_YEAR are skipped, as the nightly run would drop them anyway
  * safe to run twice: the second run adds nothing

Usage
  python add_plot_csv_to_merge.py [--src data/manual_lakes] [--out data] [--drive] [--dry-run]
  --drive   pull the two merged files from Drive first and push them back afterwards
"""
import argparse
import glob
import os
import re
import sys

import pandas as pd

START_YEAR = 2016
WSE_COL = 'Water Surface Elevation - values(m)'

# lake key (from the file name plot_<key>.csv) -> station settings.
# station_id / lat / lon are taken from the merged file when a Hydroweb station of that name is already in it;
# the values below are only the fallback. Albert has no Hydroweb id in this repo: the id below is a manual
# placeholder (outside the real Hydroweb range). Replace it with the real Hydroweb id if you know it, before
# the first run, because changing it later would leave the first rows under the old id.
LAKES = {
    'L_victoria': {'name': 'Victoria', 'station_id': 1300000000016, 'lat': -1.0, 'lon': 33.0},
    'L_albert':   {'name': 'Albert',   'station_id': 9100000000001, 'lat': 1.683, 'lon': 30.917},
}


def read_plot_csv(path, key):
    d = pd.read_csv(path)
    if d.shape[1] < 2:
        raise SystemExit(f'{path}: expected at least 2 columns')
    cols = list(d.columns)
    date_c = cols[0]
    val_c = next((c for c in cols if 'values' in c.lower()), cols[1])
    unc_c = next((c for c in cols if 'uncertainty' in c.lower()), None)
    out = pd.DataFrame({
        'date_dt': pd.to_datetime(d[date_c], format='mixed', errors='coerce'),
        'wse': pd.to_numeric(d[val_c], errors='coerce'),
        'unc': pd.to_numeric(d[unc_c], errors='coerce') if unc_c else float('nan'),
    }).dropna(subset=['date_dt', 'wse'])
    out['date'] = out['date_dt'].dt.strftime('%Y-%m-%d')
    out = out.sort_values(['date', 'unc'], na_position='last').drop_duplicates('date', keep='first')
    return out.sort_values('date').reset_index(drop=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'manual_lakes'))
    ap.add_argument('--out', default=os.environ.get('ALTIMETRY_OUT_DIR', 'data'))
    ap.add_argument('--drive', action='store_true')
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()

    csv_path = os.path.join(a.out, 'merged_altimetry_stations.csv')
    xlsx_path = os.path.join(a.out, 'merged_altimetry_stations.xlsx')
    names = ['merged_altimetry_stations.csv', 'merged_altimetry_stations.xlsx']
    svc = folder = None
    if a.drive:
        import drive_sync
        svc = drive_sync.drive_service()
        folder = os.environ['DRIVE_FOLDER_ID'].strip()
        os.makedirs(a.out, exist_ok=True)
        drive_sync.pull(svc, folder, a.out, names, [])
    if not os.path.exists(csv_path):
        sys.exit(f'{csv_path} not found. Run with --drive, or point --out at the folder that holds it.')

    header = pd.read_csv(csv_path, nrows=0).columns.tolist()
    need = ['source', 'station_id', 'type', 'location/river_name', 'latitude', 'longitude', 'date', WSE_COL,
            'uncertainty (m)', 'station_uid']
    miss = [c for c in need if c not in header]
    if miss:
        sys.exit(f'merged file lacks columns {miss}; run the nightly update first.')
    hist = pd.read_csv(csv_path, usecols=['source', 'station_id', 'type', 'location/river_name', 'latitude',
                                          'longitude', 'date', 'station_uid'])
    hist['date'] = hist['date'].astype(str)

    files = sorted(glob.glob(os.path.join(a.src, 'plot_L_*.csv')))
    if not files:
        sys.exit(f'no plot_L_*.csv in {a.src}')
    add = []
    for path in files:
        key = re.sub(r'^plot_|\.csv$', '', os.path.basename(path))
        cfg = LAKES.get(key)
        if cfg is None:
            print(f'{os.path.basename(path)}: no entry for {key} in LAKES, skipped. Add one at the top of this script.')
            continue
        sid, lat, lon, label = cfg['station_id'], cfg['lat'], cfg['lon'], f"{cfg['name']} Lake"
        # reuse the id / position already in the merged file for this lake
        same = hist[(hist['source'].str.lower() == 'hydroweb') & (hist['type'] == 'lake')
                    & (hist['station_id'].astype('int64') == sid)]
        if same.empty:
            same = hist[(hist['source'].str.lower() == 'hydroweb')
                        & hist['location/river_name'].astype(str).str.lower().str.contains(cfg['name'].lower())]
        if len(same):
            sid = int(same['station_id'].iloc[0])
            lat, lon = float(same['latitude'].iloc[0]), float(same['longitude'].iloc[0])
            label = str(same['location/river_name'].iloc[0])
            print(f'{key}: using station {sid} ({label}) at {lat}, {lon} from the merged file')
        else:
            print(f'{key}: NOT in the merged file yet -> new station {sid} at {lat}, {lon} (fallback values)')
        uid = f'hydroweb:{sid}'
        d = read_plot_csv(path, key)
        d = d[d['date_dt'].dt.year >= START_YEAR]
        have = set(hist.loc[hist['station_uid'] == uid, 'date'])
        new = d[~d['date'].isin(have)]
        print(f'  {len(d)} days in the file from {START_YEAR}, {len(d) - len(new)} already merged, '
              f'{len(new)} to add ({new["date"].min() if len(new) else "-"} to {new["date"].max() if len(new) else "-"})')
        if new.empty:
            continue
        add.append(pd.DataFrame({
            'source': 'Hydroweb', 'station_id': sid, 'type': 'lake', 'location/river_name': label,
            'latitude': lat, 'longitude': lon, 'date': new['date'].values, WSE_COL: new['wse'].values,
            'uncertainty (m)': new['unc'].values, 'station_uid': uid}))

    if not add:
        print('Nothing to add.')
        return
    rows = pd.concat(add, ignore_index=True).reindex(columns=header)
    print(f'Total new rows: {len(rows)}')
    if a.dry_run:
        print(rows.head(3).to_string())
        print('Dry run: nothing written.')
        return

    rows.to_csv(csv_path, mode='a', header=False, index=False)
    print(f'Appended to {csv_path}')
    if os.path.exists(xlsx_path):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(xlsx_path)
            ws = wb['Merged']
            xhdr = [c.value for c in ws[1]]
            for r in rows.reindex(columns=xhdr).itertuples(index=False):
                ws.append([None if pd.isna(v) else (v.item() if hasattr(v, 'item') else v) for v in r])
            wb.save(xlsx_path)
            print(f'Appended to {xlsx_path}')
        except Exception as e:
            print(f'xlsx not updated ({type(e).__name__}: {e}); the next nightly run rebuilds it from the CSV.')
    if a.drive:
        import drive_sync
        drive_sync.push(svc, folder, a.out, names)
    print('Done. The next nightly run fills the derived columns and climate for these rows.')


if __name__ == '__main__':
    main()
